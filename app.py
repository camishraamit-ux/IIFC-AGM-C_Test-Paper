import random
import urllib.parse

import docx
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from fpdf import FPDF

st.set_page_config(
    page_title="IIFCL AGM MCQ Assessment",
    page_icon="🎯",
    layout="centered",
)

st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;600;700;800;900&display=swap');

html, body, [class*="css"] {
    font-family: 'Inter', sans-serif;
}

.block-container {
    padding-top: 0.6rem !important;
    max-width: 720px !important;
}

.reg-hero {
    background: linear-gradient(135deg, #1e3a8a, #2563eb);
    padding: 40px;
    border-radius: 20px;
    text-align: center;
    color: white;
    margin-bottom: 25px;
}

.top-info-bar {
    display: flex;
    justify-content: space-between;
    padding: 12px 18px;
    background: linear-gradient(135deg, #1e3a8a, #2563eb);
    color: white;
    border-radius: 14px;
    margin-bottom: 15px;
    font-weight: 600;
}

.question-card {
    border: 2px solid #1e3a8a;
    border-radius: 16px;
    padding: 24px;
    margin-bottom: 20px;
    background: white;
}

.question-text {
    font-size: 24px;
    font-weight: 800;
}

.review-card {
    padding: 18px;
    border-radius: 14px;
    margin-bottom: 15px;
}

.correct {
    background: #ecfdf5;
    border-left: 8px solid #10b981;
}

.incorrect {
    background: #fef2f2;
    border-left: 8px solid #ef4444;
}

.unanswered {
    background: #fffbeb;
    border-left: 8px solid #f59e0b;
}
</style>
""",
    unsafe_allow_html=True,
)


# --------------------------------------------------------------------------
# Data extraction
# --------------------------------------------------------------------------
@st.cache_data
def extract_mcqs_from_docx(docx_path):
    """Reads a docx file containing one table with question rows
    (No | Question | Opt A | Opt B | Opt C | Opt D) and another set of
    rows giving the correct answer (No | <blank> | Answer text), and
    returns a DataFrame of merged, scoreable questions."""

    questions_dict = {}
    answers_dict = {}

    doc = docx.Document(docx_path)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]

            if len(cells) < 3:
                continue

            if not cells[0].isdigit():
                continue

            q_no = int(cells[0])

            if len(cells) >= 6:
                questions_dict[q_no] = {
                    "id": q_no,
                    "question": cells[1],
                    "options": cells[2:6],
                }
            elif len(cells) == 3:
                answers_dict[q_no] = cells[2]

    records = []

    for q_no, q_data in questions_dict.items():
        ans_text = answers_dict.get(q_no, "").strip()

        if not ans_text:
            continue

        correct_index = None

        for idx, opt in enumerate(q_data["options"]):
            if opt.strip().lower() == ans_text.strip().lower():
                correct_index = idx
                break

        if correct_index is None:
            # Answer text didn't match any option exactly -- skip rather
            # than silently score it wrong for everyone.
            continue

        records.append(
            {
                "id": q_data["id"],
                "question": q_data["question"],
                "options": q_data["options"],
                "correct_answer_text": ans_text,
                "correct_index": correct_index,
            }
        )

    return pd.DataFrame(records)


# --------------------------------------------------------------------------
# PDF scorecard
# --------------------------------------------------------------------------
class ScorecardPDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 10, "IIFCL AGM ASSESSMENT REPORT", ln=True, align="C")
        self.ln(2)


def generate_pdf_scorecard(
    candidate_name,
    subject_name,
    score,
    total_questions,
    percentage,
    detailed_report,
):
    pdf = ScorecardPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "", 10)

    pdf.cell(0, 8, f"Candidate : {candidate_name}", ln=True)
    pdf.cell(0, 8, f"Subject : {subject_name}", ln=True)
    pdf.cell(0, 8, f"Score : {score}/{total_questions}", ln=True)
    pdf.cell(0, 8, f"Percentage : {percentage}%", ln=True)
    pdf.ln(5)

    for item in detailed_report:
        pdf.multi_cell(0, 5, f"{item['q_num']}. {item['question']}")
        pdf.cell(0, 5, f"Your Answer : {item['user_answer'] or '(not answered)'}", ln=True)
        pdf.cell(0, 5, f"Correct Answer : {item['correct_answer']}", ln=True)
        pdf.ln(3)

    return bytes(pdf.output(dest="S"))


# --------------------------------------------------------------------------
# Config & session state
# --------------------------------------------------------------------------
SUBJECT_FILES = {
    "IIFCL AGM Test 1": "IIFCL AGM Test 1.docx",
    "IIFCL AGM Test 2": "IIFCL AGM Test 2.docx",
}

defaults = {
    "test_started": False,
    "candidate_name": "",
    "selected_subject": "",
    "assessment_df": None,
    "user_answers": {},
    "submitted": False,
    "current_q_index": 0,
}

for k, v in defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v


def restart():
    st.session_state.test_started = False
    st.session_state.submitted = False
    st.session_state.current_q_index = 0
    st.session_state.user_answers = {}
    st.session_state.assessment_df = None


# --------------------------------------------------------------------------
# Screen 1: Registration
# --------------------------------------------------------------------------
if not st.session_state.test_started:
    st.markdown(
        """
        <div class="reg-hero">
            <h1>🎯 IIFCL AGM Assessment</h1>
            <p>Start your test below</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    with st.form("registration"):
        name = st.text_input("Candidate Name")
        subject = st.selectbox("Select Paper", list(SUBJECT_FILES.keys()))
        q_count = st.selectbox("Questions", [10, 20, 25, 30, 50])
        start = st.form_submit_button("Start Assessment")

    if start:
        if not name.strip():
            st.error("Please enter your name to begin.")
            st.stop()

        try:
            df = extract_mcqs_from_docx(SUBJECT_FILES[subject])
        except FileNotFoundError:
            st.error(
                f"Question file '{SUBJECT_FILES[subject]}' was not found. "
                "Make sure it's uploaded alongside app.py in the repo."
            )
            st.stop()

        if df.empty:
            st.error("No valid questions could be loaded from that paper.")
            st.stop()

        st.session_state.assessment_df = df.sample(
            n=min(q_count, len(df)),
            random_state=random.randint(1, 100000),
        ).reset_index(drop=True)

        st.session_state.candidate_name = name.strip()
        st.session_state.selected_subject = subject
        st.session_state.user_answers = {}
        st.session_state.current_q_index = 0
        st.session_state.submitted = False
        st.session_state.test_started = True
        st.rerun()

# --------------------------------------------------------------------------
# Screen 2: Question navigation
# --------------------------------------------------------------------------
elif st.session_state.test_started and not st.session_state.submitted:
    df_test = st.session_state.assessment_df
    total_q = len(df_test)
    curr_idx = st.session_state.current_q_index
    row = df_test.iloc[curr_idx]
    valid_options = row["options"]

    st.markdown(
        f"""
        <div class="top-info-bar">
            <span>👤 {st.session_state.candidate_name}</span>
            <span>Question {curr_idx + 1} / {total_q}</span>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.progress((curr_idx + 1) / total_q)

    st.markdown(
        f"""
        <div class="question-card">
            <div class="question-text">{row['question']}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    saved_choice = st.session_state.user_answers.get(curr_idx)

    choice = st.radio(
        "Answer",
        options=list(range(len(valid_options))),
        format_func=lambda x: f"{chr(65 + x)}. {valid_options[x]}",
        key=f"q_{curr_idx}",
        index=saved_choice,
        label_visibility="collapsed",
    )

    if choice is not None:
        st.session_state.user_answers[curr_idx] = choice

    col1, col2, col3 = st.columns([1, 1, 1])

    with col1:
        if curr_idx > 0:
            if st.button("⬅️ Previous", use_container_width=True):
                st.session_state.current_q_index -= 1
                st.rerun()

    with col2:
        st.caption(f"Answered: {len(st.session_state.user_answers)}/{total_q}")

    with col3:
        if curr_idx < total_q - 1:
            if st.button("Next ➡️", use_container_width=True):
                st.session_state.current_q_index += 1
                st.rerun()
        else:
            if st.button("✅ Submit Test", use_container_width=True, type="primary"):
                st.session_state.submitted = True
                st.rerun()

# --------------------------------------------------------------------------
# Screen 3: Results
# --------------------------------------------------------------------------
elif st.session_state.submitted:
    df_test = st.session_state.assessment_df

    score = 0
    detailed_report = []

    for idx, row in df_test.iterrows():
        user_choice = st.session_state.user_answers.get(idx)
        correct_idx = row["correct_index"]

        status = "unanswered"
        user_answer = ""

        if user_choice is not None:
            user_answer = row["options"][user_choice]
            if user_choice == correct_idx:
                score += 1
                status = "correct"
            else:
                status = "incorrect"

        detailed_report.append(
            {
                "q_num": f"Q{idx + 1}",
                "question": row["question"],
                "user_answer": user_answer,
                "correct_answer": row["correct_answer_text"],
                "status_type": status,
            }
        )

    incorrect_count = sum(1 for x in detailed_report if x["status_type"] == "incorrect")
    unanswered_count = sum(1 for x in detailed_report if x["status_type"] == "unanswered")
    percentage = round(score / len(df_test) * 100, 2)

    st.markdown(
        f"""
        <div class="reg-hero">
            <h1>🎉 Assessment Complete</h1>
            <p>{st.session_state.candidate_name} — {st.session_state.selected_subject}</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    fig = go.Figure()
    fig.add_trace(
        go.Pie(
            labels=["Correct", "Incorrect", "Unanswered"],
            values=[score, incorrect_count, unanswered_count],
            hole=0.70,
            marker=dict(colors=["#10b981", "#ef4444", "#f59e0b"]),
        )
    )
    fig.update_layout(
        height=350,
        annotations=[dict(text=f"{percentage}%", x=0.5, y=0.5, showarrow=False, font=dict(size=28))],
        showlegend=True,
    )
    st.plotly_chart(fig, use_container_width=True)

    m1, m2, m3 = st.columns(3)
    m1.metric("Score", f"{score}/{len(df_test)}")
    m2.metric("Correct", score)
    m3.metric("Incorrect", incorrect_count)

    with st.expander("📋 Detailed Review"):
        for item in detailed_report:
            css_class = item["status_type"]
            st.markdown(
                f"""
                <div class="review-card {css_class}">
                    <b>{item['q_num']}. {item['question']}</b><br>
                    Your Answer: {item['user_answer'] or '(not answered)'}<br>
                    Correct Answer: {item['correct_answer']}
                </div>
                """,
                unsafe_allow_html=True,
            )

    share_text = (
        f"🎯 IIFCL AGM Assessment\n\n"
        f"Candidate: {st.session_state.candidate_name}\n"
        f"Subject: {st.session_state.selected_subject}\n"
        f"Score: {score}/{len(df_test)}\n"
        f"Percentage: {percentage}%"
    )
    whatsapp_url = "https://api.whatsapp.com/send?text=" + urllib.parse.quote(share_text)

    st.markdown(
        f"""
        <a href="{whatsapp_url}" target="_blank" style="text-decoration:none;">
            <button style="
                width:100%;
                padding:16px;
                background:#25D366;
                color:white;
                border:none;
                border-radius:12px;
                font-size:18px;
                font-weight:bold;
                cursor:pointer;">
                📲 Share on WhatsApp
            </button>
        </a>
        """,
        unsafe_allow_html=True,
    )

    st.write("")

    pdf_data = generate_pdf_scorecard(
        st.session_state.candidate_name,
        st.session_state.selected_subject,
        score,
        len(df_test),
        percentage,
        detailed_report,
    )

    st.download_button(
        "📄 Download PDF Scorecard",
        pdf_data,
        file_name=f"{st.session_state.candidate_name}_Scorecard.pdf",
        mime="application/pdf",
        use_container_width=True,
    )

    if st.button("🔄 Take Another Test", use_container_width=True):
        restart()
        st.rerun()
